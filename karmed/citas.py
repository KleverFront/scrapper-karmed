"""Fetch and export appointments (citas) from the Karmed derma agenda.

The agenda uses FullCalendar; its event source is modules/bdd.php, which returns
JSON for a date range + doctor (medico) + branch (sucursal). Unlike the patient
list, this is already JSON, so no HTML scraping is needed.
"""
from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from karmed.client import KarmedClient

AGENDA_PAGE = "modules/agenda/"
CITAS_ENDPOINT = "modules/bdd.php"  # NOT modules/agenda/bdd.php — that path 404s/returns the page


@dataclass(slots=True)
class Cita:
    id: str
    fecha: str
    hora: str
    paciente: str
    num_hist: str
    estado: str          # raw color code from the agenda (status encoding)
    medico: str
    start: str
    end: str


@dataclass(slots=True)
class Medico:
    id: str
    nombre: str


def get_agenda_context(client: KarmedClient) -> tuple[list[Medico], str]:
    """Read the doctors dropdown and the branch id from the agenda page."""
    html = client.get(AGENDA_PAGE).text
    medicos: list[Medico] = []
    sel = re.search(r"<select[^>]*id=['\"]medicos['\"][^>]*>(.*?)</select>", html, re.S | re.I)
    if sel:
        for val, txt in re.findall(r"<option[^>]*value=['\"]?([^'\">]+)['\"]?[^>]*>([^<]*)", sel.group(1), re.I):
            medicos.append(Medico(id=val.strip(), nombre=re.sub(r"\s+", " ", txt).strip()))
    suc = re.search(r"id=['\"]cod_sucursal['\"][^>]*value=['\"]([^'\"]*)", html, re.I)
    return medicos, (suc.group(1) if suc else "0")


def _split_dt(value: str) -> tuple[str, str]:
    """'YYYY-MM-DD HH:MM:SS' -> ('YYYY-MM-DD', 'HH:MM')."""
    value = (value or "").strip()
    if " " in value:
        d, t = value.split(" ", 1)
        return d, t[:5]
    return value, ""


def fetch_citas(
    client: KarmedClient,
    medico: Medico,
    sucursal: str,
    start: str,
    end: str,
) -> list[Cita]:
    """Fetch citas for one doctor in a datetime range (inclusive endpoints)."""
    resp = client.post(
        CITAS_ENDPOINT,
        data={"start": start, "end": end, "index": medico.id, "cod_sucursal": sucursal},
    )
    try:
        events = json.loads(resp.text)
    except json.JSONDecodeError:
        return []  # endpoint returned the page shell (invalid params) -> no data
    citas: list[Cita] = []
    for ev in events:
        fecha, hora = _split_dt(ev.get("start", ""))
        citas.append(
            Cita(
                id=str(ev.get("id", "")),
                fecha=fecha,
                hora=hora,
                paciente=re.sub(r"\s+", " ", str(ev.get("title", ""))).strip(),
                num_hist=str(ev.get("num_hist", "")),
                estado=str(ev.get("color", "")),
                medico=medico.nombre,
                start=str(ev.get("start", "")),
                end=str(ev.get("end", "")),
            )
        )
    return citas


def fetch_all_citas(
    client: KarmedClient,
    medicos: list[Medico],
    sucursal: str,
    year_from: int = 2015,
    year_to: int | None = None,
    progress=None,
) -> list[Cita]:
    """Fetch every cita across the given years for all doctors, deduped by id."""
    year_to = year_to or (datetime.now().year + 1)
    by_id: dict[str, Cita] = {}
    years = list(range(year_from, year_to + 1))
    total_steps = max(1, len(medicos) * len(years))
    step = 0
    for medico in medicos:
        for year in years:
            start = f"{year}-01-01 00:00:00"
            end = f"{year}-12-31 23:59:59"
            for cita in fetch_citas(client, medico, sucursal, start, end):
                by_id[cita.id] = cita
            step += 1
            if progress:
                progress(step, total_steps, len(by_id))
    return sorted(by_id.values(), key=lambda c: (c.start, c.id))


_HEADERS = ["Fecha", "Hora", "Paciente", "N° Historia", "Estado (color)", "Médico"]


def export_citas_to_excel(citas: list[Cita], out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Citas"

    header_fill = PatternFill("solid", fgColor="8957E5")
    header_font = Font(bold=True, color="FFFFFF")
    for col, title in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for c in citas:
        v = asdict(c)
        ws.append([v["fecha"], v["hora"], v["paciente"], v["num_hist"], v["estado"], v["medico"]])

    widths = [12, 8, 40, 14, 16, 24]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}{len(citas) + 1}"
    wb.save(out_path)
    return out_path


def default_citas_name() -> str:
    return f"citas_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
