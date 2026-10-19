"""Fetch and parse the patient list from the Karmed derma app, and export it to Excel."""
from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from karmed.client import KarmedClient

# consultaP.php returns the whole list; an empty name still matches everything,
# but "a" is the safe catch-all the UI uses.
LIST_ENDPOINT = "modules/consultaP.php"

_ROW_RE = re.compile(r"<tr\b[^>]*>(.*?)</tr>", re.I | re.S)
_CELL_RE = re.compile(r"<td\b[^>]*>(.*?)</td>", re.I | re.S)
_TAG_RE = re.compile(r"<[^>]+>")
_HIST_RE = re.compile(r"ir_a_historia\((\d+)\)")
_WS_RE = re.compile(r"\s+")


@dataclass(slots=True)
class Patient:
    historia: str
    nombre: str
    cedula: str
    edad: str
    telefonos: str


def _text(cell_html: str) -> str:
    return _WS_RE.sub(" ", _TAG_RE.sub("", cell_html)).strip()


def parse_patients(html: str) -> list[Patient]:
    """Extract patient rows from a consultaP.php response.

    Each data row has six cells (historia, paciente, cedula, edad, telefonos,
    actions); the header row inside <thead> uses <th>, so it is skipped here.
    """
    patients: list[Patient] = []
    for row in _ROW_RE.findall(html):
        cells = _CELL_RE.findall(row)
        if len(cells) < 5:
            continue  # header row or layout row, no <td> data
        texts = [_text(c) for c in cells]
        hist = texts[0]
        m = _HIST_RE.search(row)
        if m:
            hist = m.group(1)
        patients.append(
            Patient(
                historia=hist,
                nombre=texts[1],
                cedula=texts[2],
                edad=texts[3],
                telefonos=texts[4],
            )
        )
    return patients


def fetch_patients(client: KarmedClient, cliente: str = "3", query: str = "a") -> list[Patient]:
    resp = client.get(LIST_ENDPOINT, params={"nombre": query, "cliente": cliente})
    return parse_patients(resp.text)


_HEADERS = ["No. Historia", "Paciente", "Cédula", "Edad", "Teléfonos"]


def export_to_excel(patients: list[Patient], out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pacientes"

    header_fill = PatternFill("solid", fgColor="1F6FEB")
    header_font = Font(bold=True, color="FFFFFF")
    for col, title in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for patient in patients:
        values = asdict(patient)
        ws.append([values["historia"], values["nombre"], values["cedula"], values["edad"], values["telefonos"]])

    widths = [14, 36, 16, 20, 18]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}{len(patients) + 1}"

    wb.save(out_path)
    return out_path


def default_export_name() -> str:
    return f"pacientes_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def save_cache(patients: list[Patient], path: str | Path) -> str:
    """Persist the scraped list to disk so reopening is instant (no re-fetch).

    Returns the saved_at timestamp so callers don't need to re-read the file.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    saved_at = datetime.now().isoformat(timespec="seconds")
    payload = {"saved_at": saved_at, "patients": [asdict(p) for p in patients]}
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)  # atomic: never leave a half-written cache
    return saved_at


def load_cache(path: str | Path) -> tuple[list[Patient], str | None]:
    """Load a previously saved list. Returns ([], None) when no valid cache exists."""
    path = Path(path)
    if not path.is_file():
        return [], None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        patients = [Patient(**row) for row in data.get("patients", [])]
        return patients, data.get("saved_at")
    except (json.JSONDecodeError, TypeError, ValueError):
        return [], None


def diff_new(previous: list[Patient], current: list[Patient]) -> list[str]:
    """Historia numbers present in current but not in previous (newly added patients)."""
    old = {p.historia for p in previous}
    return [p.historia for p in current if p.historia not in old]
